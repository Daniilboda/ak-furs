"""Проверка шага 7: канонические имена менеджеров."""

from __future__ import annotations

import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from crm_dedupe import dedupe_order_lines, read_crm
from crm_normalize import (
    apply_manager_normalization,
    load_manager_aliases,
    normalize_manager,
)

DOCS_CRM = Path(__file__).resolve().parents[2] / "docs" / "01_crm_orders_export.csv"
DOCS_CATALOG = Path(__file__).resolve().parents[2] / "docs" / "03_catalog_prices_reference.xlsx"


def test_reference_variants_map_to_canonical():
    aliases = load_manager_aliases(DOCS_CATALOG)
    assert normalize_manager("вероника", aliases) == "Вероника Белова"
    assert normalize_manager(" В. Белова", aliases) == "Вероника Белова"
    assert normalize_manager("Veronika B", aliases) == "Вероника Белова"
    assert normalize_manager("АЛИНА С.", aliases) == "Алина Сергеева"
    assert normalize_manager("alina.s", aliases) == "Алина Сергеева"
    assert normalize_manager("P. Orlova", aliases) == "Полина Орлова"
    assert normalize_manager("полина орлова ", aliases) == "Полина Орлова"
    assert normalize_manager("marina_k", aliases) == "Марина Котова"
    assert normalize_manager("М.Котова", aliases) == "Марина Котова"
    assert normalize_manager("olga.e", aliases) == "Ольга Ершова"
    assert normalize_manager("О. Ершова", aliases) == "Ольга Ершова"


def test_every_variant_from_sheet_resolves():
    aliases = load_manager_aliases(DOCS_CATALOG)
    from openpyxl import load_workbook

    sheet = load_workbook(DOCS_CATALOG, data_only=True)["Менеджеры"]
    for row in list(sheet.iter_rows(values_only=True))[1:]:
        canonical = str(row[0]).strip()
        for variant in str(row[1]).split("|"):
            assert normalize_manager(variant, aliases) == canonical
        assert normalize_manager(canonical, aliases) == canonical


def test_unknown_manager_raises():
    aliases = load_manager_aliases(DOCS_CATALOG)
    with pytest.raises(ValueError, match="не найден"):
        normalize_manager("Неизвестный", aliases)


def test_apply_manager_normalization_does_not_mutate_input():
    aliases = load_manager_aliases(DOCS_CATALOG)
    original = [{"Менеджер": "olga.e"}]
    updated = apply_manager_normalization(original, aliases)
    assert original[0]["Менеджер"] == "olga.e"
    assert updated[0]["Менеджер"] == "Ольга Ершова"


def test_current_export_managers_become_canonical():
    aliases = load_manager_aliases(DOCS_CATALOG)
    records = apply_manager_normalization(dedupe_order_lines(read_crm(DOCS_CRM)), aliases)
    allowed = set(aliases.values())
    assert {record["Менеджер"] for record in records} <= allowed


def test_load_manager_aliases_from_substituted_file(tmp_path: Path):
    path = tmp_path / "catalog.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Менеджеры"
    sheet.append(["Каноническое имя", "Варианты написания в CRM", "Дата выхода", "Дата увольнения"])
    sheet.append(["Тест Тестов", "тест | t.test", None, None])
    workbook.save(path)
    aliases = load_manager_aliases(path)
    assert normalize_manager("t.test", aliases) == "Тест Тестов"
