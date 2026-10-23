"""Проверка шага 9: тот же код на подменённой выгрузке CRM."""

from __future__ import annotations

import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from crm_dedupe import CRM_HEADER
from run_crm_cleanup import cleaned_output_name, main, run_cleanup

DOCS_CATALOG = Path(__file__).resolve().parents[2] / "docs" / "03_catalog_prices_reference.xlsx"


def _write_crm(path: Path, rows: list[list[str]]) -> None:
    lines = [";".join(CRM_HEADER)]
    lines.extend(";".join(row) for row in rows)
    path.write_text("\n".join(lines), encoding="utf-8-sig")


def _row(
    order: str,
    line: str,
    *,
    exported: str = "2026-10-01 09:00",
    status: str = "PAID",
    price: str = "10 000",
    qty: str = "1",
    discount: str = "",
    manager: str = "olga.e",
) -> list[str]:
    values = {column: "" for column in CRM_HEADER}
    values.update(
        {
            "Номер заказа": order,
            "№ строки": line,
            "Дата создания": "2026-01-01",
            "Клиент": "CL-1",
            "Менеджер": manager,
            "Источник": "Сайт",
            "Артикул": "SOB-SH-01",
            "Кол-во": qty,
            "Цена за ед.": price,
            "Валюта": "руб",
            "Скидка": discount,
            "Статус": status,
            "Выгружено": exported,
        }
    )
    return [values[column] for column in CRM_HEADER]


def test_run_cleanup_on_substituted_crm_counts_q1(tmp_path: Path):
    crm_path = tmp_path / "new_crm.csv"
    _write_crm(
        crm_path,
        [
            _row("A-00075", "1", exported="2026-09-01 10:00", price="100"),
            _row("a-75", "1", exported="2026-10-01 09:00", price="200"),  # дубль, останется эта
            _row("A-00075", "2", price="300"),
            _row("A-00100", "1", price="400"),
        ],
    )
    records, (orders, lines) = run_cleanup(crm_path, DOCS_CATALOG, only="dedupe")
    assert orders == 2
    assert lines == 3
    assert len(records) == 3
    kept_line_1 = next(
        record for record in records if record["№ строки"].strip() == "1" and "75" in record["Номер заказа"]
    )
    assert kept_line_1["Цена за ед."] == "200.00"


def test_full_cleanup_writes_normalized_file_for_new_data(tmp_path: Path):
    crm_path = tmp_path / "interview_crm.csv"
    out_dir = tmp_path / "normalized_crm"
    _write_crm(
        crm_path,
        [
            _row("A-1", "1", status="оплачено", discount="10%", manager="вероника"),
            _row("A-1", "1", status="PAID", discount="10%", manager="вероника", exported="2026-10-02 12:00"),
        ],
    )
    exit_code = main(
        [
            "--crm",
            str(crm_path),
            "--catalog",
            str(DOCS_CATALOG),
            "--out-dir",
            str(out_dir),
            "--only",
            "all",
        ]
    )
    assert exit_code == 0
    assert cleaned_output_name(crm_path) == "interview_crm_cleaned.csv"
    out_file = out_dir / "interview_crm_cleaned.csv"
    assert out_file.is_file()
    text = out_file.read_text(encoding="utf-8-sig")
    assert "Оплачен" in text
    assert "Вероника Белова" in text
    assert "Цена за ед. после скидки" in text
    assert ";1;" in text or text.splitlines()[1].startswith("1;")
    assert "A-1" not in text


def test_catalog_workbook_keeps_all_sheets_after_cleanup(tmp_path: Path):
    catalog_copy = tmp_path / "03_catalog_prices_reference.xlsx"
    shutil.copy2(DOCS_CATALOG, catalog_copy)
    sheets_before = load_workbook(catalog_copy).sheetnames
    crm_path = tmp_path / "01_crm_orders_export.csv"
    _write_crm(crm_path, [_row("A-1", "1")])
    main(
        [
            "--crm",
            str(crm_path),
            "--catalog",
            str(catalog_copy),
            "--out-dir",
            str(tmp_path / "out"),
        ]
    )
    sheets_after = load_workbook(catalog_copy).sheetnames
    assert sheets_after == sheets_before
    assert (tmp_path / "out" / "01_crm_orders_export_cleaned.csv").is_file()


def test_only_dedupe_skips_manager_normalization(tmp_path: Path):
    crm_path = tmp_path / "crm.csv"
    _write_crm(crm_path, [_row("A-1", "1", manager="olga.e")])
    records, _ = run_cleanup(crm_path, DOCS_CATALOG, only="dedupe")
    assert records[0]["Менеджер"] == "olga.e"
    assert "Цена за ед. после скидки" not in records[0]


def test_only_normalize_keeps_duplicate_rows(tmp_path: Path):
    crm_path = tmp_path / "crm.csv"
    _write_crm(
        crm_path,
        [
            _row("A-1", "1", exported="2026-09-01 10:00"),
            _row("A-1", "1", exported="2026-10-01 09:00"),
        ],
    )
    records, (orders, lines) = run_cleanup(crm_path, DOCS_CATALOG, only="normalize")
    assert len(records) == 2
    assert orders == 1
    assert lines == 1
    assert all(record["Менеджер"] == "Ольга Ершова" for record in records)
