"""Нормализация полей CRM: статусы, скидки, даты, цены, источники и менеджеры."""

from __future__ import annotations

import re
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Literal

from openpyxl import load_workbook

STATUS_PAID = "Оплачен"
STATUS_CANCELLED = "Отменён"
STATUS_REFUND = "Возврат"
STATUS_IN_PROGRESS = "В работе"

_PAID = frozenset({"paid", "оплачен", "оплачено"})
_CANCELLED = frozenset({"cancelled", "отменен"})
_REFUND = frozenset({"refund", "возврат"})


def normalize_status(value: str) -> str:
    """Свести статус к одному из четырёх: Оплачен, Отменён, Возврат, В работе.

    Всё, что не подходит под первые три, считается «В работе».
    Регистр и пробелы по краям не учитываются. «ё» и «е» считаются одним знаком.
    """
    key = value.strip().lower().replace("ё", "е")
    if key in _PAID:
        return STATUS_PAID
    if key in _CANCELLED:
        return STATUS_CANCELLED
    if key in _REFUND:
        return STATUS_REFUND
    return STATUS_IN_PROGRESS


def apply_status_normalization(records: list[dict[str, str]]) -> list[dict[str, str]]:
    """Вернуть копии записей со сведённым статусом. Исходный список не меняется."""
    result: list[dict[str, str]] = []
    for record in records:
        updated = dict(record)
        updated["Статус"] = normalize_status(record["Статус"])
        result.append(updated)
    return result


def parse_money(value: str | float | int) -> float:
    """Разобрать сумму в рублях: пробелы, запятая/точка, запись с «млн»."""
    if isinstance(value, (int, float)):
        return float(value)

    text = value.strip().replace("\xa0", " ").replace(" ", "")
    if text == "":
        raise ValueError(f"пустая сумма: {value!r}")

    lower = text.lower()
    million = False
    if "млн" in lower:
        million = True
        text = re.sub(r"млн\.?", "", lower, flags=re.IGNORECASE).strip()

    if "," in text and "." in text:
        if text.rfind(".") > text.rfind(","):
            text = text.replace(",", "")
        else:
            text = text.replace(".", "").replace(",", ".")
    elif "," in text:
        text = text.replace(",", ".")

    try:
        amount = float(text)
    except ValueError as error:
        raise ValueError(f"не удалось разобрать сумму: {value!r}") from error
    if million:
        amount *= 1_000_000
    return amount


def format_money(value: str | float | int) -> str:
    """Записать сумму с точкой и копейками: 50000.00."""
    return f"{parse_money(value):.2f}"


def apply_canonical_unit_prices(records: list[dict]) -> list[dict]:
    """Записать цены за единицу в формате с копейками. Исходный список не меняется."""
    result: list[dict] = []
    for record in records:
        updated = dict(record)
        updated["Цена за ед."] = format_money(record["Цена за ед."])
        if "Цена за ед. после скидки" in record:
            updated["Цена за ед. после скидки"] = format_money(
                record["Цена за ед. после скидки"]
            )
        result.append(updated)
    return result


SOURCE_INSTAGRAM = "Instagram"
SOURCE_TELEGRAM = "Telegram"
SOURCE_REFERRAL = "Рекомендация"
SOURCE_YANDEX_MAPS = "Яндекс Карты"
SOURCE_SITE = "Сайт"

_SOURCE_ALIASES = {
    "ig": SOURCE_INSTAGRAM,
    "insta": SOURCE_INSTAGRAM,
    "инстаграм": SOURCE_INSTAGRAM,
    "instagram": SOURCE_INSTAGRAM,
    "telegram": SOURCE_TELEGRAM,
    "телеграм": SOURCE_TELEGRAM,
    "tg": SOURCE_TELEGRAM,
    "рекомендация": SOURCE_REFERRAL,
    "реком.": SOURCE_REFERRAL,
    "сарафан": SOURCE_REFERRAL,
    "яндекс карты": SOURCE_YANDEX_MAPS,
    "я.карты": SOURCE_YANDEX_MAPS,
    "yandex maps": SOURCE_YANDEX_MAPS,
    "сайт": SOURCE_SITE,
    "заявка с сайта": SOURCE_SITE,
    "site": SOURCE_SITE,
}


def normalize_source(value: str) -> str:
    """Свести источник к каноническому имени."""
    key = value.strip().casefold()
    if key not in _SOURCE_ALIASES:
        raise ValueError(f"источник не найден среди известных: {value!r}")
    return _SOURCE_ALIASES[key]


def apply_source_normalization(records: list[dict]) -> list[dict]:
    """Записать канонический источник. Исходный список не меняется."""
    result: list[dict] = []
    for record in records:
        updated = dict(record)
        updated["Источник"] = normalize_source(str(record["Источник"]))
        result.append(updated)
    return result


CURRENCY_RUB = "руб."

_CURRENCY_ALIASES = {
    "rub": CURRENCY_RUB,
    "rur": CURRENCY_RUB,
    "руб": CURRENCY_RUB,
    "руб.": CURRENCY_RUB,
    "₽": CURRENCY_RUB,
}


def normalize_currency(value: str) -> str:
    """Свести валюту к «руб.»."""
    key = value.strip().casefold()
    if key not in _CURRENCY_ALIASES:
        raise ValueError(f"валюта не найдена среди известных: {value!r}")
    return _CURRENCY_ALIASES[key]


def apply_currency_normalization(records: list[dict]) -> list[dict]:
    """Записать валюту как «руб.». Исходный список не меняется."""
    result: list[dict] = []
    for record in records:
        updated = dict(record)
        updated["Валюта"] = normalize_currency(str(record["Валюта"]))
        result.append(updated)
    return result


DiscountKind = Literal["percent", "amount"]


def parse_discount(value: str) -> tuple[DiscountKind, float]:
    """Разобрать скидку.

    «N%» — процент от суммы строки. Число без % — рубли на всю строку.
    Пустое поле — скидка 0 рублей.
    """
    text = value.strip()
    if text == "":
        return "amount", 0.0
    if text.endswith("%"):
        percent = parse_money(text[:-1])
        return "percent", percent
    return "amount", parse_money(text)


def unit_price_after_discount(price: str, quantity: str, discount: str) -> float:
    """Цена за единицу после скидки.

    Процент считается от суммы строки (цена × количество).
    Сумма без % вычитается из суммы строки целиком.
    """
    unit = parse_money(price)
    qty = parse_money(quantity)
    if qty == 0:
        raise ValueError("количество не может быть нулём")
    line_sum = unit * qty
    kind, amount = parse_discount(discount)
    if kind == "percent":
        after = line_sum * (1 - amount / 100)
    else:
        after = line_sum - amount
    return after / qty


def apply_discount_fields(
    records: list[dict[str, str]],
) -> list[dict[str, str | float]]:
    """Вернуть копии записей с полем «Цена за ед. после скидки». Исходный список не меняется."""
    result: list[dict[str, str | float]] = []
    for record in records:
        updated: dict[str, str | float] = dict(record)
        updated["Цена за ед. после скидки"] = unit_price_after_discount(
            record["Цена за ед."],
            record["Кол-во"],
            record["Скидка"],
        )
        result.append(updated)
    return result


def load_manager_aliases(path: Path | str) -> dict[str, str]:
    """Прочитать лист «Менеджеры»: ключ — вариант написания в нижнем регистре, значение — канон.

    Файл справочника только читается и не сохраняется: остальные листы не затрагиваются.
    """
    workbook = load_workbook(path, data_only=True)
    if "Менеджеры" not in workbook.sheetnames:
        raise ValueError(f"{path}: нет листа «Менеджеры»")
    sheet = workbook["Менеджеры"]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"{path}: лист «Менеджеры» пуст")

    aliases: dict[str, str] = {}
    for row in rows[1:]:
        if not row or row[0] is None:
            continue
        canonical = str(row[0]).strip()
        variants_cell = "" if row[1] is None else str(row[1])
        names = [canonical, *[part.strip() for part in variants_cell.split("|")]]
        for name in names:
            if name == "":
                continue
            key = name.casefold()
            previous = aliases.get(key)
            if previous is not None and previous != canonical:
                raise ValueError(
                    f"вариант {name!r} относится и к {previous!r}, и к {canonical!r}"
                )
            aliases[key] = canonical
    return aliases


def normalize_manager(value: str, aliases: dict[str, str]) -> str:
    """Свести написание менеджера к каноническому имени из справочника."""
    key = value.strip().casefold()
    if key not in aliases:
        raise ValueError(f"менеджер не найден в справочнике: {value!r}")
    return aliases[key]


def apply_manager_normalization(
    records: list[dict[str, str]],
    aliases: dict[str, str],
) -> list[dict[str, str]]:
    """Вернуть копии записей с каноническим именем менеджера. Исходный список не меняется."""
    result: list[dict[str, str]] = []
    for record in records:
        updated = dict(record)
        updated["Менеджер"] = normalize_manager(record["Менеджер"], aliases)
        result.append(updated)
    return result


_RU_MONTHS = {
    "янв": 1,
    "фев": 2,
    "мар": 3,
    "апр": 4,
    "май": 5,
    "мая": 5,
    "июн": 6,
    "июл": 7,
    "авг": 8,
    "сен": 9,
    "окт": 10,
    "ноя": 11,
    "дек": 12,
}

_EXCEL_EPOCH = date(1899, 12, 30)
_RU_TEXT_DATE = re.compile(
    r"^\s*(\d{1,2})\s+([а-яёА-ЯЁ]+)\s+(\d{4})\s*$",
    flags=re.IGNORECASE,
)


def parse_creation_date(value: str) -> date:
    """Разобрать «Дата создания» из разных форматов выгрузки."""
    text = value.strip()
    if text == "":
        raise ValueError("пустая дата создания")

    if re.fullmatch(r"\d{5}", text):
        return _EXCEL_EPOCH + timedelta(days=int(text))

    for fmt in ("%Y-%m-%d", "%d.%m.%Y", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass

    match = _RU_TEXT_DATE.match(text)
    if match is not None:
        day = int(match.group(1))
        month_key = match.group(2).casefold()[:3]
        year = int(match.group(3))
        month = _RU_MONTHS.get(month_key)
        if month is None:
            raise ValueError(f"неизвестный месяц в дате: {value!r}")
        return date(year, month, day)

    raise ValueError(f"не удалось разобрать дату создания: {value!r}")


def format_creation_date(value: str) -> str:
    """Записать дату создания как ДД.ММ.ГГГГ, например 03.03.2026."""
    return parse_creation_date(value).strftime("%d.%m.%Y")


def apply_canonical_creation_dates(records: list[dict]) -> list[dict]:
    """Записать «Дата создания» в формате ДД.ММ.ГГГГ. Исходный список не меняется."""
    result: list[dict] = []
    for record in records:
        updated = dict(record)
        updated["Дата создания"] = format_creation_date(str(record["Дата создания"]))
        result.append(updated)
    return result

