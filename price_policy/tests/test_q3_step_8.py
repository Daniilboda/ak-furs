"""Шаг 8 этапа 3: Q3 на подменённых заказах и каталоге."""

from __future__ import annotations

import csv
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
sys.path.insert(0, str(Path(__file__).resolve().parents[2] / "crm_cleanup"))

from catalog_prices import (
    PRICE_HISTORY_HEADER,
    PRICE_HISTORY_SHEET,
    RENAMES_HEADER,
    RENAMES_SHEET,
    price_history_sheet_is_hidden,
)
from crm_dedupe import CRM_HEADER
from q3_summary import compute_q3_from_cleaned, compute_q3_from_raw
from run_crm_cleanup import cleaned_output_name, run_cleanup, write_crm_csv
from run_price_policy import main

DOCS_CATALOG = Path(__file__).resolve().parents[2] / "docs" / "03_catalog_prices_reference.xlsx"
DOCS_CRM = Path(__file__).resolve().parents[2] / "docs" / "01_crm_orders_export.csv"
CLEANED_CRM = (
    Path(__file__).resolve().parents[2]
    / "crm_cleanup"
    / "normalized_crm"
    / "01_crm_orders_export_cleaned.csv"
)


def _write_catalog(path: Path) -> None:
    wb = Workbook()
    price = wb.active
    price.title = PRICE_HISTORY_SHEET
    price.append(list(PRICE_HISTORY_HEADER))
    price.append(["ART-1", date(2025, 1, 1), None, 100000, 90000])
    price.sheet_state = "hidden"
    ren = wb.create_sheet(RENAMES_SHEET)
    ren.append(list(RENAMES_HEADER))
    ren.append(["OLD-1", "ART-1", date(2025, 3, 1), "старый может быть после перехода"])
    managers = wb.create_sheet("Менеджеры")
    managers.append(
        ["Каноническое имя", "Варианты написания в CRM", "Дата выхода", "Дата увольнения"]
    )
    managers.append(["Вероника Белова", "Вероника Белова | вероника", date(2025, 1, 1), None])
    wb.save(path)


def _crm_row(order: str, *, status: str, price: str, discount: str, article: str) -> list[str]:
    values = {column: "" for column in CRM_HEADER}
    values.update(
        {
            "Номер заказа": order,
            "№ строки": "1",
            "Дата создания": "2025-11-05",
            "Клиент": "CL-1",
            "Менеджер": "вероника",
            "Источник": "Сайт",
            "Артикул": article,
            "Кол-во": "1",
            "Цена за ед.": price,
            "Валюта": "руб",
            "Скидка": discount,
            "Статус": status,
            "Выгружено": "2026-10-01 09:00",
        }
    )
    return [values[column] for column in CRM_HEADER]


def _write_crm(path: Path, rows: list[list[str]]) -> None:
    lines = [";".join(CRM_HEADER)]
    lines.extend(";".join(row) for row in rows)
    path.write_text("\n".join(lines), encoding="utf-8-sig")


def test_q3_on_substituted_crm_and_hidden_price_sheet(tmp_path: Path):
    catalog = tmp_path / "interview_catalog.xlsx"
    raw_crm = tmp_path / "interview_crm.csv"
    _write_catalog(catalog)
    assert price_history_sheet_is_hidden(catalog) is True

    _write_crm(
        raw_crm,
        [
            # 100000 - 20% = 80000 < 90000 → violation
            _crm_row("A-1", status="PAID", price="100000", discount="20%", article="OLD-1"),
            # отменён ниже минимума → не считаем
            _crm_row("A-2", status="CANCELLED", price="100000", discount="20%", article="ART-1"),
            # равно минимуму → не нарушение
            _crm_row("A-3", status="NEW", price="90000", discount="", article="ART-1"),
        ],
    )
    result = compute_q3_from_raw(raw_crm, catalog)
    assert result["q3_price_violation_lines_total"] == 1
    assert result["q3_price_violation_lines_by_manager"] == {"Вероника Белова": 1}


def test_substituted_differs_from_current_docs():
    current = compute_q3_from_cleaned(CLEANED_CRM, DOCS_CATALOG)
    assert current["q3_price_violation_lines_total"] == 68
    assert current["q3_price_violation_lines_total"] != 1


def test_cli_writes_results_json_and_csv(tmp_path: Path, capsys):
    catalog = tmp_path / "catalog.xlsx"
    raw_crm = tmp_path / "crm.csv"
    out_csv = tmp_path / "q3_price_violations.csv"
    _write_catalog(catalog)
    _write_crm(
        raw_crm,
        [_crm_row("A-1", status="PAID", price="100000", discount="20%", article="OLD-1")],
    )
    cleaned, _ = run_cleanup(raw_crm, catalog)
    cleaned_path = tmp_path / cleaned_output_name(raw_crm)
    write_crm_csv(cleaned_path, cleaned)

    code = main(
        [
            "--crm",
            str(cleaned_path),
            "--catalog",
            str(catalog),
            "--out",
            str(out_csv),
        ]
    )
    assert code == 0
    stdout = capsys.readouterr().out
    payload = json.loads(stdout.split("saved=")[0])
    assert payload["q3_price_violation_lines_total"] == 1
    assert out_csv.is_file()
    text = out_csv.read_text(encoding="utf-8-sig")
    assert "q3_price_violation_lines_total" in text
    assert "Вероника Белова" in text
