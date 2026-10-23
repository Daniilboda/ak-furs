"""Шаг 1 этапа 3: прайс и переименования."""

from __future__ import annotations

import shutil
import sys
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from catalog_prices import (
    PRICE_HISTORY_HEADER,
    PRICE_HISTORY_SHEET,
    RENAMES_HEADER,
    RENAMES_SHEET,
    price_history_sheet_is_hidden,
    read_article_renames,
    read_price_history,
)

DOCS_CATALOG = Path(__file__).resolve().parents[2] / "docs" / "03_catalog_prices_reference.xlsx"


def test_current_catalog_has_hidden_price_history_with_required_columns():
    assert price_history_sheet_is_hidden(DOCS_CATALOG) is True
    rows = read_price_history(DOCS_CATALOG)
    assert rows
    assert set(rows[0]) >= {
        "Артикул",
        "Действует с",
        "Действует по",
        "Минимально допустимая цена, ₽",
    }
    assert rows[0]["Артикул"]
    assert isinstance(rows[0]["Действует с"], date)
    assert isinstance(rows[0]["Минимально допустимая цена, ₽"], float)


def test_current_catalog_renames_have_old_new_and_date():
    rows = read_article_renames(DOCS_CATALOG)
    assert rows
    assert set(rows[0]) == set(RENAMES_HEADER)
    assert rows[0]["Старый артикул"]
    assert rows[0]["Новый артикул"]
    assert isinstance(rows[0]["Дата перехода"], date)


def test_open_ended_price_period_allowed():
    rows = read_price_history(DOCS_CATALOG)
    open_ended = [row for row in rows if row["Действует по"] is None]
    assert open_ended


def test_read_price_history_from_substituted_workbook(tmp_path: Path):
    path = tmp_path / "catalog.xlsx"
    wb = Workbook()
    ws = wb.active
    ws.title = PRICE_HISTORY_SHEET
    ws.append(list(PRICE_HISTORY_HEADER))
    ws.append(["ART-1", date(2025, 1, 1), date(2025, 6, 30), 100000, 90000])
    ws.sheet_state = "hidden"
    ren = wb.create_sheet(RENAMES_SHEET)
    ren.append(list(RENAMES_HEADER))
    ren.append(["OLD-1", "ART-1", date(2025, 3, 1), "тест"])
    wb.save(path)

    assert price_history_sheet_is_hidden(path) is True
    history = read_price_history(path)
    assert len(history) == 1
    assert history[0]["Артикул"] == "ART-1"
    assert history[0]["Минимально допустимая цена, ₽"] == 90000.0
    renames = read_article_renames(path)
    assert renames[0]["Старый артикул"] == "OLD-1"
    assert renames[0]["Новый артикул"] == "ART-1"


def test_missing_price_sheet_raises(tmp_path: Path):
    path = tmp_path / "bad.xlsx"
    shutil.copy2(DOCS_CATALOG, path)
    wb = __import__("openpyxl").load_workbook(path)
    # удалить скрытый лист через создание книги без него
    wb2 = Workbook()
    wb2.active.title = "Каталог"
    wb2.save(path)
    with pytest.raises(ValueError, match="Прайс_история"):
        read_price_history(path)
